# -*- coding: utf-8 -*-
"""Generate DOCX/PDF documentation from the canonical Markdown sources.

Markdown under docs/ is the source of truth. This script deliberately reads the
same files a human reviews instead of maintaining a second hardcoded API model.
Run with a Python that has python-docx installed, for example:

    "C:/Program Files/LibreOffice/program/python.exe" docs/gen_docs.py

PDF is produced with LibreOffice headless when soffice is available.
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
SOFFICE = Path(r"C:/Program Files/LibreOffice/program/soffice.exe")

# (language, document base name, Markdown source)
BOOKS = [
    ("en", "Standard", "STANDARD.md"),
    ("en", "API", "API.md"),
    ("en", "Platforms", "PLATFORMS.md"),
    ("ru", "Standard", "STANDARD.ru.md"),
    ("ru", "API", "API.ru.md"),
]


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("__", "")
    return text.strip()


def add_inline(paragraph, text: str) -> None:
    # Keep the parser intentionally small: bold spans are the only inline style
    # that survives, because it is the one used by the standard documents.
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(clean_inline(part))
            run.bold = True
        else:
            paragraph.add_run(clean_inline(part))


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Light Grid Accent 1"
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            cell.text = clean_inline(row[col_index]) if col_index < len(row) else ""
            if row_index == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph()


def render_markdown(doc, text: str) -> None:
    lines = text.splitlines()
    code: list[str] = []
    in_code = False
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph()
                run = paragraph.add_run("\n".join(code))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                paragraph.paragraph_format.left_indent = Pt(12)
                code.clear()
                in_code = False
            else:
                flush_table()
                in_code = True
            continue
        if in_code:
            code.append(line)
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            if re.match(r"^\|?[\s:|-]+\|?$", stripped):
                continue
            table_rows.append([part.strip() for part in stripped.strip("|").split("|")])
            continue
        flush_table()

        if not stripped:
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            doc.add_heading(clean_inline(heading.group(2)), level=min(level, 4))
            continue
        if stripped.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), stripped[2:])
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            add_inline(doc.add_paragraph(style="List Number"), numbered.group(1))
            continue
        add_inline(doc.add_paragraph(), stripped)

    flush_table()


def build_book(lang: str, base: str, source: str) -> Path | None:
    path = DOCS / source
    if not path.exists():
        print(f"skip missing source: {path}")
        return None

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    text = path.read_text(encoding="utf-8")
    render_markdown(doc, text)

    out_dir = DOCS / "docx" / lang
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"FemboyEngine_{base}_{lang.upper()}.docx"
    doc.save(out)
    print(f"docx: {out}")
    return out


def convert_pdf(docx_path: Path | None, lang: str) -> None:
    if docx_path is None or not SOFFICE.exists():
        return
    pdf_dir = DOCS / "pdf" / lang
    pdf_dir.mkdir(parents=True, exist_ok=True)
    profile = ROOT / "docs" / ".lo_profile"
    command = [
        str(SOFFICE),
        "--headless",
        f"-env:UserInstallation=file:///{profile.as_posix()}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(pdf_dir),
        str(docx_path),
    ]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode != 0:
        print(result.stdout)
        print(result.stderr, file=sys.stderr)


def main() -> int:
    # The generated tree is an artifact; rebuild it from scratch so stale
    # per-OS or old camelCase documents cannot survive a regeneration.
    for directory in (DOCS / "docx", DOCS / "pdf"):
        if directory.exists():
            shutil.rmtree(directory)
    for lang, base, source in BOOKS:
        docx_path = build_book(lang, base, source)
        convert_pdf(docx_path, lang)
    profile = DOCS / ".lo_profile"
    if profile.exists():
        shutil.rmtree(profile, ignore_errors=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
